"""engine.annotator 的回归测试。

历史上的致命缺陷：``engine/annotator.py`` 用命名空间 URI 拼 XML 标签时漏掉
Clark 记法的花括号（写成 ``f"{CT_NS}Override"`` 而非 ``f"{{{CT_NS}}}Override"``），
导致 lxml 抛出 ``ValueError: Invalid tag name 'http://...Override'``。
而 ``word/comments.xml`` 只有在文档已被 Word 写过批注时才存在，所以这个分支
对**任何没有现成批注的文档**都会触发——即对绝大多数用户文档批注功能完全不可用。
同时 ``cli.cmd_annotate`` 把一切失败统一报成"仅支持 .docx 文件"，把真实根因掩盖了。

本文件用"无 comments.xml 的最小 docx"固定这条路径，防止再次回归。
"""
import io
import os
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from docx import Document
from lxml import etree

from engine import annotator

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
PR_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
COMMENTS_CT = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"
)

SAMPLE_DOCX = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "tests", "fixtures", "sample.docx",
)


def _make_docx(path, paragraphs=("标题行", "正文第一段", "正文第二段", "结束段")):
    """生成一个没有任何批注的 docx（即无 word/comments.xml）。"""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))
    with zipfile.ZipFile(str(path)) as z:
        assert "word/comments.xml" not in z.namelist()
    return str(path)


def _errors(items, start=0):
    """构造 RuleEngine.validate() 风格的错误列表。"""
    return [
        {"type": etype, "description": desc, "paragraph_index": start + i}
        for i, (desc, etype) in enumerate(items)
    ]


@pytest.fixture
def tmp_docx(tmp_path):
    return _make_docx(tmp_path / "src.docx")


class TestAnnotateNewDocument:
    """核心回归：无现成批注的文档必须能成功写入批注。"""

    def test_basic_annotate_succeeds(self, tmp_path, tmp_docx):
        out = tmp_path / "out.docx"
        errs = _errors([("标题字号应为二号", "format"),
                        ("正文应首行缩进 2 字符", "format")])
        assert annotator.annotate_document(tmp_docx, str(out), errs) is True
        assert annotator.LAST_ERROR == ""

    def test_produces_comment_parts(self, tmp_path, tmp_docx):
        out = tmp_path / "out.docx"
        errs = _errors([("标题字号应为二号", "format"),
                        ("正文应首行缩进 2 字符", "format"),
                        ("参考文献至少 5 条", "count")])
        annotator.annotate_document(tmp_docx, str(out), errs)

        with zipfile.ZipFile(str(out)) as z:
            names = z.namelist()
            # 批注正文必须落盘
            assert "word/comments.xml" in names
            comments = etree.fromstring(z.read("word/comments.xml"))
            doc = etree.fromstring(z.read("word/document.xml"))
            ct = etree.fromstring(z.read("[Content_Types].xml"))
            rels = etree.fromstring(z.read("word/_rels/document.xml.rels"))

        assert len(comments.findall(f"{{{W_NS}}}comment")) == 3
        # document.xml 内每条评论都必须有三段引用标记
        assert len(doc.findall(f".//{{{W_NS}}}commentRangeStart")) == 3
        assert len(doc.findall(f".//{{{W_NS}}}commentRangeEnd")) == 3
        assert len(doc.findall(f".//{{{W_NS}}}commentReference")) == 3

        # 新建文档时必须补注册 content type 与 relationship
        overrides = ct.findall(f"{{{CT_NS}}}Override")
        assert any(
            o.get("ContentType") == COMMENTS_CT and o.get("PartName") == "/word/comments.xml"
            for o in overrides
        )
        rels_list = rels.findall(f"{{{PR_NS}}}Relationship")
        assert any(r.get("Target") == "comments.xml" for r in rels_list)

    def test_output_remains_openable(self, tmp_path, tmp_docx):
        """产物必须是 python-docx 能正常打开的合法 docx。"""
        out = tmp_path / "out.docx"
        annotator.annotate_document(
            tmp_docx, str(out), _errors([("需要修复", "format")] * 5)
        )
        doc = Document(str(out))
        assert len(doc.paragraphs) == 4

    def test_comment_id_unique_and_sequential(self, tmp_path, tmp_docx):
        out = tmp_path / "out.docx"
        annotator.annotate_document(tmp_docx, str(out), _errors(
            [(f"错误 {i}", "format") for i in range(7)]
        ))
        with zipfile.ZipFile(str(out)) as z:
            comments = etree.fromstring(z.read("word/comments.xml"))
        ids = [int(c.get(f"{{{W_NS}}}id")) for c in comments.findall(f"{{{W_NS}}}comment")]
        assert ids == sorted(ids)
        assert len(set(ids)) == len(ids)

    def test_max_per_paragraph_caps_comments(self, tmp_path, tmp_docx):
        # 同一段落上报 10 条错误，最多只能写入 max_per_paragraph 条
        errs = [{"type": "format", "description": f"错误 {i}", "paragraph_index": 0}
                for i in range(10)]
        out = tmp_path / "out.docx"
        annotator.annotate_document(tmp_docx, str(out), errs, max_per_paragraph=3)
        with zipfile.ZipFile(str(out)) as z:
            comments = etree.fromstring(z.read("word/comments.xml"))
        assert len(comments.findall(f"{{{W_NS}}}comment")) == 3

    def test_out_of_range_paragraph_index_is_skipped(self, tmp_path, tmp_docx):
        out = tmp_path / "out.docx"
        errs = [{"type": "format", "description": "越界", "paragraph_index": 99},
                {"type": "format", "description": "合法", "paragraph_index": 0}]
        annotator.annotate_document(tmp_docx, str(out), errs)
        with zipfile.ZipFile(str(out)) as z:
            comments = etree.fromstring(z.read("word/comments.xml"))
        assert len(comments.findall(f"{{{W_NS}}}comment")) == 1

    def test_multiple_paragraphs_all_get_comments(self, tmp_path, tmp_docx):
        out = tmp_path / "out.docx"
        errs = _errors([("标题问题", "format"), ("正文问题", "format"),
                        ("段末问题", "format")], start=0)
        annotator.annotate_document(tmp_docx, str(out), errs)
        with zipfile.ZipFile(str(out)) as z:
            comments = etree.fromstring(z.read("word/comments.xml"))
        assert len(comments.findall(f"{{{W_NS}}}comment")) == 3


class TestAnnotateExistingComments:
    """文档已有批注时：复用既有 comments.xml，不重复注册 content type。"""

    def test_appends_without_duplicate_registration(self, tmp_path, tmp_docx):
        out = tmp_path / "step1.docx"
        annotator.annotate_document(
            tmp_docx, str(out), _errors([("第一轮", "format")])
        )
        with zipfile.ZipFile(str(out)) as z:
            names = z.namelist()
        assert "word/comments.xml" in names

        final = tmp_path / "step2.docx"
        annotator.annotate_document(str(out), str(final), _errors([("第二轮", "format")]))

        with zipfile.ZipFile(str(final)) as z:
            comments = etree.fromstring(z.read("word/comments.xml"))
            ct = etree.fromstring(z.read("[Content_Types].xml"))
        # 两轮各 1 条，共 2 条
        assert len(comments.findall(f"{{{W_NS}}}comment")) == 2
        # content type 只注册一次
        assert sum(
            1 for o in ct.findall(f"{{{CT_NS}}}Override")
            if o.get("ContentType") == COMMENTS_CT
        ) == 1


class TestAnnotateNoOp:
    """没有可定位错误时不应崩溃、也不应写入批注。"""

    def test_no_annotatable_errors_returns_true(self, tmp_path, tmp_docx):
        # paragraph_index 为 None（结构级错误）时无法定位段落
        errs = [{"type": "count", "description": "参考文献不足", "paragraph_index": None},
                {"type": "format", "description": "", "paragraph_index": 0}]
        out = tmp_path / "out.docx"
        assert annotator.annotate_document(tmp_docx, str(out), errs) is True
        with zipfile.ZipFile(str(out)) as z:
            assert "word/comments.xml" not in z.namelist()

    def test_empty_errors_returns_true(self, tmp_path, tmp_docx):
        out = tmp_path / "out.docx"
        assert annotator.annotate_document(tmp_docx, str(out), []) is True


class TestAnnotateUnsupported:
    """不支持的格式必须返回 False 并给出真实原因。"""

    @pytest.mark.parametrize("suffix", [".doc", ".pdf"])
    def test_unsupported_format(self, tmp_path, suffix):
        src = tmp_path / f"file{suffix}"
        src.write_bytes(b"not a real document")
        out = tmp_path / "out.docx"
        assert annotator.annotate_document(str(src), str(out), _errors([("x", "format")])) is False
        assert annotator.LAST_ERROR == "仅支持 .docx 文件"

    def test_last_error_reset_on_success(self, tmp_path, tmp_docx):
        annotator.LAST_ERROR = "上一次遗留的失败原因"
        out = tmp_path / "out.docx"
        annotator.annotate_document(tmp_docx, str(out), _errors([("x", "format")]))
        assert annotator.LAST_ERROR == ""


@pytest.mark.skipif(not os.path.exists(SAMPLE_DOCX), reason="sample docx not found")
def test_real_word_document_without_comments(tmp_path):
    """用真实 Word 生成的样本（无 comments.xml）跑通全流程。"""
    with zipfile.ZipFile(SAMPLE_DOCX) as z:
        assert "word/comments.xml" not in z.namelist()

    out = tmp_path / "sample_annotated.docx"
    errs = _errors([("论文标题字号应为二号(22pt)", "format"),
                    ("论文标题应居中", "format")])
    assert annotator.annotate_document(SAMPLE_DOCX, str(out), errs) is True

    with zipfile.ZipFile(str(out)) as z:
        comments = etree.fromstring(z.read("word/comments.xml"))
        doc = etree.fromstring(z.read("word/document.xml"))
    assert len(comments.findall(f"{{{W_NS}}}comment")) == 2
    assert len(doc.findall(f".//{{{W_NS}}}commentReference")) == 2
    Document(str(out))  # python-docx 可正常打开


class TestStdoutHardening:
    """回归：GBK 控制台下 print_good/print_bad/print_info 不得抛 UnicodeEncodeError。

    中文 Windows 控制台默认编码是 GBK，而 ui/printers.py 用的 ✓ ✗ ℹ 在 GBK 中
    不存在，未加固时任何一条状态消息都会把 CLI 打死（连错误处理链一起崩）。
    """

    def test_harden_stdout_relaxes_errors(self, monkeypatch):
        import ui.printers as printers

        buf = io.BytesIO()
        fake = io.TextIOWrapper(buf, encoding="gbk", errors="strict", line_buffering=True)
        monkeypatch.setattr(printers.sys, "stdout", fake)

        printers._harden_stdout()
        assert fake.errors == "replace"
        # 加固后 GBK 无法编码的符号降级为替换字符，而不是抛异常
        fake.write("✓ ✗ ℹ\n")
        fake.flush()
        assert fake.errors == "replace"

    def test_harden_stdout_keeps_original_encoding(self, monkeypatch):
        """只放宽 errors，不改编码——GBK 能覆盖全部中文，改编码会导致乱码。"""
        import ui.printers as printers

        fake = io.TextIOWrapper(io.BytesIO(), encoding="gbk", errors="strict")
        monkeypatch.setattr(printers.sys, "stdout", fake)
        printers._harden_stdout()
        assert fake.encoding.lower().startswith("gbk") or fake.encoding == "cp936"
        assert fake.errors == "replace"

    def test_harden_stdout_ignores_non_reconfigurable_streams(self, monkeypatch):
        import ui.printers as printers

        monkeypatch.setattr(printers.sys, "stdout", io.StringIO())
        # StringIO 没有 reconfigure，必须静默跳过而不是抛异常
        printers._harden_stdout()
