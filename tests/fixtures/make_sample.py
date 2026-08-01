"""生成测试用样本 docx。运行：python tests/fixtures/make_sample.py"""
import os
from docx import Document
from docx.shared import Pt


def make_sample(path):
    doc = Document()
    p1 = doc.add_paragraph("测试标题")
    p1.style = doc.styles["Heading 1"]
    for r in p1.runs:
        r.font.size = Pt(18)
        r.font.name = "黑体"
    p2 = doc.add_paragraph("这是正文内容，用于测试格式化引擎。")
    for r in p2.runs:
        r.font.size = Pt(12)
        r.font.name = "宋体"
    p3 = doc.add_paragraph("第二章 测试章节")
    p3.style = doc.styles["Heading 1"]
    for r in p3.runs:
        r.font.size = Pt(16)
    p4 = doc.add_paragraph("正文段落二。")
    for r in p4.runs:
        r.font.size = Pt(12)
    doc.save(path)
    print(f"CREATED {path}")


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    make_sample(os.path.join(here, "sample.docx"))
