"""数字编号标题层级 + 结构性标题识别的回归测试。

修复前有两个缺陷：

1. ``RE_HEADING_H3 = ^\\d+\\s*[\\.．]`` 把 "1."、"1.1"、"1.1.1" 全部匹配成
   三级标题，所以 "1.1 课题研究背景" 与 "1.1.1 课题研究背景" 都是 Lv3，
   而 "1. 数据来源" 也是 Lv3——层级信息完全丢失。
2. ``摘  要``、``参  考  文  献`` 这类中间带空格的结构性标题不被任何正则识别，
   被当成普通正文套上了 2 字符首行缩进。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import logging

import pytest
from docx import Document

from format import Config
from format.ops import heading
from format.ops._paragraph_helpers import get_outline_level
from format.ops.heading import HeadingOp
from format.ops.paragraph_classifier import ParagraphClassifier


# ---------------------------------------------------------------------------
# 数字编号层级
# ---------------------------------------------------------------------------

class TestNumericHeadingLevel:
    def test_depth_comes_from_dot_segments(self):
        assert heading.numeric_heading_level("1. 数据来源") == 1
        assert heading.numeric_heading_level("1．数据来源") == 1
        assert heading.numeric_heading_level("1.1 课题研究背景及意义") == 2
        assert heading.numeric_heading_level("1.1.1 课题研究背景") == 3
        assert heading.numeric_heading_level("1.1.1.1 更细粒度") == 4

    def test_depth_capped_at_four(self):
        assert heading.numeric_heading_level("1.1.1.1.1 更深") == 4

    def test_leading_whitespace_ignored(self):
        assert heading.numeric_heading_level("   2.1 标题") == 2

    def test_dates_and_versions_are_not_headings(self):
        # 首段超过两位数字视为年份/版本号，不是章节号
        assert heading.numeric_heading_level("2024.5 数据") is None
        assert heading.numeric_heading_level("1998.12 会议") is None

    def test_oversize_text_is_not_a_heading(self):
        assert heading.numeric_heading_level(
            "1.1 " + "正文内容" * 40
        ) is None

    def test_plain_text_is_none(self):
        assert heading.numeric_heading_level("普通正文段落") is None
        assert heading.numeric_heading_level("") is None
        assert heading.numeric_heading_level(None) is None


class TestHeadingOpDetectType:
    """HeadingOp.detect_type 必须给出正确的层级而不是全部 h3。"""

    @pytest.fixture
    def op(self):
        return HeadingOp(Config("official"))

    @pytest.mark.parametrize("text,expected", [
        ("1. 数据来源", "h1"),
        ("1.1 课题研究背景及意义", "h2"),
        ("1.1.1 课题研究背景", "h3"),
        ("1.1.1.1 实验设备", "h4"),
        ("2. 相关工作", "h1"),
        ("二、结论", "h1"),
        ("（一）研究目标", "h2"),
        ("(1) 设备清单", "h4"),
        ("普通正文段落", "body"),
    ])
    def test_level_detection(self, op, text, expected):
        assert op.detect_type(text) == expected

    def test_attachment_still_wins(self, op):
        # RE_ATTACHMENT 只匹配裸附件标识（^...$），带内容的附件行不在此列
        assert op.detect_type("附件1：") == "attachment"
        assert op.detect_type("附件：") == "attachment"


# ---------------------------------------------------------------------------
# 结构性标题（摘要 / 参考文献 / 致谢 ...）
# ---------------------------------------------------------------------------

class TestStructuralHeadingKind:
    def test_recognized_markers(self):
        assert heading.structural_heading_kind("摘要") == "h1"
        assert heading.structural_heading_kind("ABSTRACT") == "h1"
        assert heading.structural_heading_kind("Abstract") == "h1"
        for text in ("参考文献", "致谢", "附录", "目录", "结论", "REFERENCES"):
            assert heading.structural_heading_kind(text) == "h1", text

    def test_ideographic_space_padding(self):
        # 缺陷场景：文档在这些词上留视觉空格
        assert heading.structural_heading_kind("摘  要") == "h1"
        assert heading.structural_heading_kind("参  考  文  献") == "h1"
        assert heading.structural_heading_kind("致　谢") == "h1"
        assert heading.structural_heading_kind("摘  要：") == "h1"

    def test_keywords_marker(self):
        assert heading.structural_heading_kind("关键词：神经网络 深度学习") == "keywords"
        assert heading.structural_heading_kind("KEYWORDS:  deep learning") == "keywords"

    def test_marker_followed_by_content_is_not_a_heading(self):
        # 参考文献条目，不是章节标题
        assert heading.structural_heading_kind("参考文献[1]张三 论文") is None
        assert heading.structural_heading_kind("附录A 实验数据说明") is None
        assert heading.structural_heading_kind("结论部分讨论了局限性。") is None

    def test_ordinary_text_is_none(self):
        assert heading.structural_heading_kind("普通正文段落") is None
        assert heading.structural_heading_kind("摘要方法综述") is None
        assert heading.structural_heading_kind("") is None
        assert heading.structural_heading_kind(None) is None


class TestClassifierOutlineLevels:
    """端到端：apply_para_format 后段落的大纲级别要跟着编号深度走。

    ``w:outlineLvl`` 存的是 0-based 值，``get_outline_level`` 直接返回原值，
    所以要 +1 才是 Word 里看到的 1-based 层级。
    """

    @pytest.fixture
    def doc(self, tmp_path):
        d = Document()
        for text in (
            "摘  要",
            "本文研究了某问题。",
            "1. 数据来源",
            "1.1 课题研究背景及意义",
            "1.1.1 课题研究背景",
            "参考文献",
            "参考文献[1]张三. 某论文[J]. 2024.",
            "摘要方法综述正文",
        ):
            d.add_paragraph(text)
        return d

    def test_outline_levels_follow_numbering_depth(self, doc):
        cfg = Config("official")
        clf = ParagraphClassifier(cfg, logging.getLogger("t"))
        paras = doc.paragraphs
        for p in paras:
            clf.apply_para_format(p)

        def lv(p):
            v = get_outline_level(p)
            return v + 1 if v is not None else None

        assert lv(paras[0]) == 1      # 摘  要 → 结构性标题 Lv1
        assert lv(paras[2]) == 1      # 1. → Lv1
        assert lv(paras[3]) == 2      # 1.1 → Lv2
        assert lv(paras[4]) == 3      # 1.1.1 → Lv3
        assert lv(paras[5]) == 1      # 参考文献 → 结构性标题 Lv1
        # 参考文献条目与含"摘要"的普通正文不得被当成标题
        assert lv(paras[6]) is None
        assert lv(paras[7]) is None

    def test_structural_heading_uses_heading_font(self, doc):
        """缺陷场景：摘要被当成正文，套了正文字体而不是标题字体。"""
        cfg = Config("official")
        clf = ParagraphClassifier(cfg, logging.getLogger("t"))
        clf.apply_para_format(doc.paragraphs[0])
        run = doc.paragraphs[0].runs[0]
        assert run.font.name == cfg.get("h1_font")
        assert run.font.name != cfg.get("body_font")

    def test_regular_body_uses_body_font(self, doc):
        cfg = Config("official")
        clf = ParagraphClassifier(cfg, logging.getLogger("t"))
        clf.apply_para_format(doc.paragraphs[1])
        run = doc.paragraphs[1].runs[0]
        assert run.font.name == cfg.get("body_font")

    def test_contained_marker_is_treated_as_body(self, doc):
        """"参考文献[1]张三..." 是条目，不是标题，必须按正文格式。"""
        cfg = Config("official")
        clf = ParagraphClassifier(cfg, logging.getLogger("t"))
        clf.apply_para_format(doc.paragraphs[6])
        run = doc.paragraphs[6].runs[0]
        assert run.font.name == cfg.get("body_font")


# ---------------------------------------------------------------------------
# 正则仍是单一来源（供其它模块 re-import）
# ---------------------------------------------------------------------------

def test_new_regexes_exported_from_heading_module():
    for name in ("RE_HEADING_NUMERIC", "RE_HEADING_NUM_SINGLE",
                 "RE_STRUCTURAL_HEADING", "RE_KEYWORDS_LINE",
                 "strip_all_whitespace",
                 "numeric_heading_level", "structural_heading_kind"):
        assert hasattr(heading, name), f"{name} missing from format.ops.heading"


def test_strip_all_whitespace_removes_ideographic_space():
    assert heading.strip_all_whitespace("摘  要") == "摘要"
    assert heading.strip_all_whitespace("　摘　要　") == "摘要"
    assert heading.strip_all_whitespace("") == ""
    assert heading.strip_all_whitespace(None) == ""
