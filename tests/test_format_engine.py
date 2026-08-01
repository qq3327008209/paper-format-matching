import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from format import Config, FormatEngine
from format.ops.font import FontOp
from format.ops.heading import HeadingOp
from format.ops.paragraph import ParagraphOp
from format.ops.punctuation import PunctuationOp

SAMPLE_DOCX = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "tests", "fixtures", "sample.docx",
)


def test_config_presets():
    Config("official")
    Config("academic")
    Config("legal")
    assert Config("official").get("body_font") == "仿宋_GB2312"


def test_font_op():
    doc = Document()
    para = doc.add_paragraph("测试")
    run = para.runs[0]
    font_op = FontOp(Config("official"))
    font_op.set_font(run, "黑体", "Times New Roman", 16)
    assert run.font.name == "Times New Roman"
    assert run.font.size == Pt(16)


def test_paragraph_op():
    doc = Document()
    para = doc.add_paragraph("测试")
    ParagraphOp(Config("official")).set_align(para, "center")
    assert para.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_heading_op_detect():
    op = HeadingOp(Config("official"))
    assert op.detect_type("一、总则") == "h1"
    assert op.detect_type("附件1：xxx") in ("attachment", "body")
    assert op.detect_type("普通正文") in ("body", None)


def test_punctuation_op():
    result = PunctuationOp(Config("official")).fix_text("你好(hello)")
    assert "（" in result


@pytest.mark.skipif(
    not os.path.exists(SAMPLE_DOCX),
    reason="sample docx not found",
)
def test_format_engine_run(tmp_path):
    out = tmp_path / "formatted.docx"
    FormatEngine(Config("official")).run(SAMPLE_DOCX, str(out))
    assert os.path.exists(str(out))
