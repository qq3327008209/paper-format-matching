"""Word document (.docx) content extractor.

Extracts paragraphs, tables, headers, footers, and section metadata.
All numeric values are converted to conventional units:
  - page size / margins / indent: cm
  - font size, spacing: pt
  - line spacing: float (multiplier) or pt for fixed values
"""

import logging
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

logger = logging.getLogger(__name__)


class DocxExtractor:
    """Extract structured content from a .docx document via python-docx."""

    ALIGNMENT_MAP = {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    }

    LINE_SPACING_RULE_MAP = {
        WD_LINE_SPACING.SINGLE: "single",
        WD_LINE_SPACING.ONE_POINT_FIVE: "multiple",
        WD_LINE_SPACING.DOUBLE: "multiple",
        WD_LINE_SPACING.AT_LEAST: "atLeast",
        WD_LINE_SPACING.EXACTLY: "exact",
        WD_LINE_SPACING.MULTIPLE: "multiple",
    }

    def extract(self, file_path: str) -> dict:
        """Extract full document structure.

        Args:
            file_path: Path to the .docx file.

        Returns:
            Dict containing paragraphs, tables, headers, footers, sections,
            and source_format fields.
        """
        document = Document(file_path)
        paragraphs = self._extract_paragraphs(document)
        self._annotate_blank_neighbors(paragraphs)
        return {
            "paragraphs": paragraphs,
            "tables": self._extract_tables(document),
            "headers": self._extract_headers(document),
            "footers": self._extract_footers(document),
            "sections": self._extract_sections(document),
            "source_format": "docx",
        }

    @staticmethod
    def _annotate_blank_neighbors(paragraphs: list) -> None:
        """Set has_blank_before/has_blank_after on each paragraph in-place.

        A paragraph counts as blank when its stripped text is empty. The first
        paragraph has has_blank_before=False and the last has
        has_blank_after=False.
        """
        total = len(paragraphs)
        for i, item in enumerate(paragraphs):
            prev_blank = i > 0 and paragraphs[i - 1]["text"].strip() == ""
            next_blank = i < total - 1 and paragraphs[i + 1]["text"].strip() == ""
            item["has_blank_before"] = prev_blank
            item["has_blank_after"] = next_blank

    def _extract_paragraphs(self, document) -> list:
        """Extract all paragraphs with text and style attributes."""
        paragraphs: list[dict] = []
        for index, para in enumerate(document.paragraphs):
            item = {
                "index": index,
                "paragraph_index": index,
                "text": para.text,
                "style": para.style.name if para.style else None,
            }
            item.update(self._extract_run_style(para))
            item.update(self._extract_paragraph_format(para))
            paragraphs.append(item)
        return paragraphs

    # ---- run-level style ----

    def _extract_run_style(self, paragraph) -> dict:
        """Extract font-related attributes from the first meaningful run.

        Also computes paragraph-wide majority fonts (font_cn_most /
        font_en_most) by scanning all runs, plus a runs_summary histogram
        for in-paragraph consistency checks.
        """
        info: dict = {"font": None, "size": None, "bold": None, "italic": None}
        target_run = None
        for run in paragraph.runs:
            if run.text and run.text.strip():
                target_run = run
                break
        if target_run is None and paragraph.runs:
            target_run = paragraph.runs[0]
        if target_run is not None:
            font = target_run.font
            info["font"] = self._get_font_name(target_run)
            info["size"] = font.size.pt if font.size else None
            info["bold"] = font.bold
            info["italic"] = font.italic
        self._inherit_from_style_chain(paragraph, info)

        # Paragraph-wide majority fonts and run histogram (new fields).
        info["font_cn_most"], info["font_en_most"] = self._majority_fonts(paragraph)
        info["runs_summary"] = self._runs_summary(paragraph)
        return info

    @staticmethod
    def _is_cjk_char(ch: str) -> bool:
        """Return True if ch is a CJK ideograph or related punctuation."""
        cp = ord(ch)
        return (
            0x4E00 <= cp <= 0x9FFF      # CJK Unified Ideographs
            or 0x3400 <= cp <= 0x4DBF   # CJK Unified Ideographs Extension A
            or 0x20000 <= cp <= 0x2A6DF  # CJK Unified Ideographs Extension B
            or 0x3000 <= cp <= 0x303F   # CJK Symbols and Punctuation
            or 0xFF00 <= cp <= 0xFFEF   # Halfwidth and Fullwidth Forms
        )

    @staticmethod
    def _is_ascii_char(ch: str) -> bool:
        """Return True if ch is a printable ASCII character (excl. space)."""
        cp = ord(ch)
        return 0x21 <= cp <= 0x7E

    def _majority_fonts(self, paragraph) -> tuple:
        """Compute the majority eastAsia and ascii fonts across all runs.

        CJK character counts are tallied per eastAsia font, ASCII character
        counts per ascii font. The font with the largest weighted tally wins.

        Returns:
            (font_cn_most, font_en_most) -- str or None each.
        """
        cn_counter: dict = {}
        en_counter: dict = {}
        for run in paragraph.runs:
            text = run.text or ""
            cn_count = sum(1 for ch in text if self._is_cjk_char(ch))
            en_count = sum(1 for ch in text if self._is_ascii_char(ch))
            if cn_count > 0:
                ea_font = self._get_run_font_slot(run, "eastAsia", paragraph)
                if ea_font:
                    cn_counter[ea_font] = cn_counter.get(ea_font, 0) + cn_count
            if en_count > 0:
                ascii_font = self._get_run_font_slot(run, "ascii", paragraph)
                if ascii_font:
                    en_counter[ascii_font] = en_counter.get(ascii_font, 0) + en_count
        font_cn_most = max(cn_counter, key=cn_counter.get) if cn_counter else None
        font_en_most = max(en_counter, key=en_counter.get) if en_counter else None
        return font_cn_most, font_en_most

    def _runs_summary(self, paragraph) -> dict:
        """Build a histogram of font and size occurrences across all runs.

        Returns:
            {"font": {name: count, ...}, "size": {pt: count, ...}}
        """
        summary: dict = {"font": {}, "size": {}}
        for run in paragraph.runs:
            font_name = self._get_font_name(run)
            size = run.font.size.pt if run.font.size else None
            if font_name is not None:
                summary["font"][font_name] = summary["font"].get(font_name, 0) + 1
            if size is not None:
                summary["size"][size] = summary["size"].get(size, 0) + 1
        return summary

    @staticmethod
    def _get_run_font_slot(run, slot: str, paragraph=None) -> Optional[str]:
        """Read a specific font slot (eastAsia/ascii) for a run.

        Walks the run's rPr, then the paragraph style chain, then the
        document's docDefaults (styles.xml w:docDefaults). For the ascii
        slot, also falls back to python-docx's resolved run.font.name.

        In real-world .docx files, CJK fonts are frequently set only in the
        docDefaults / style chain (not on the run itself), so the eastAsia
        slot must walk the full inheritance chain including docDefaults.

        Args:
            run: The run object.
            slot: "eastAsia" or "ascii".
            paragraph: The parent paragraph for style inheritance.

        Returns:
            Font name string, or None.
        """
        slot_qn = qn(f"w:{slot}")
        # Run-level rPr.
        try:
            rPr = run._element.rPr
            if rPr is not None and rPr.rFonts is not None:
                val = rPr.rFonts.get(slot_qn)
                if val:
                    return val
        except Exception:
            pass
        # Walk paragraph style chain.
        style = paragraph.style if paragraph is not None else None
        depth = 0
        while style is not None and depth < 10:
            try:
                rPr = style.element.rPr if hasattr(style, "element") else None
                if rPr is not None and rPr.rFonts is not None:
                    val = rPr.rFonts.get(slot_qn)
                    if val:
                        return val
                style = style.base_style
            except Exception:
                break
            depth += 1
        # docDefaults fallback (styles.xml w:docDefaults/w:rPrDefault/w:rPr).
        # 真实 Word 文档中，CJK 字体经常只设在 docDefaults 里，run 自己不设。
        # 访问路径：paragraph.part.styles.element → w:docDefaults → w:rPrDefault → w:rPr → w:rFonts
        try:
            part = paragraph.part if paragraph is not None else None
            if part is not None:
                # DocumentPart.styles 是 property，返回 Styles 对象，.element 是 styles.xml 根
                styles_obj = part.styles
                if styles_obj is not None:
                    styles_element = styles_obj.element
                    doc_defaults = styles_element.find(qn("w:docDefaults"))
                    if doc_defaults is not None:
                        rpr_default = doc_defaults.find(qn("w:rPrDefault"))
                        if rpr_default is not None:
                            rpr = rpr_default.find(qn("w:rPr"))
                            if rpr is not None and rpr.rFonts is not None:
                                val = rpr.rFonts.get(slot_qn)
                                if val:
                                    return val
        except Exception:
            pass
        # For the ascii slot, fall back to python-docx's resolved name.
        if slot == "ascii":
            try:
                name = run.font.name
                if name:
                    return name
            except Exception:
                pass
        return None

    @staticmethod
    def _inherit_from_style_chain(paragraph, info: dict) -> None:
        """Fill missing fields by walking the style inheritance chain."""
        style = paragraph.style if paragraph is not None else None
        depth = 0
        while style is not None and depth < 10:
            try:
                font = style.font
                if info["size"] is None and font.size is not None:
                    info["size"] = font.size.pt
                if info["bold"] is None and font.bold is not None:
                    info["bold"] = font.bold
                if info["italic"] is None and font.italic is not None:
                    info["italic"] = font.italic
                if info["font"] is None:
                    inherited = DocxExtractor._get_style_font_name(style)
                    if inherited:
                        info["font"] = inherited
                if all(v is not None for v in (info["font"], info["size"],
                                               info["bold"], info["italic"])):
                    break
                style = style.base_style
            except Exception:
                break
            depth += 1

    @staticmethod
    def _get_style_font_name(style) -> Optional[str]:
        """Read font name from style XML (eastAsia / ascii)."""
        try:
            rPr = style.element.rPr if hasattr(style, "element") else None
            if rPr is not None and rPr.rFonts is not None:
                rFonts = rPr.rFonts
                ea = rFonts.get(qn("w:eastAsia"))
                if ea:
                    return ea
                ascii_f = rFonts.get(qn("w:ascii"))
                if ascii_f:
                    return ascii_f
        except Exception:
            pass
        try:
            return style.font.name
        except Exception:
            return None

    @staticmethod
    def _get_font_name(run) -> Optional[str]:
        """Get the full font name including CJK eastAsia font."""
        try:
            rPr = run._element.rPr
            if rPr is not None and rPr.rFonts is not None:
                rFonts = rPr.rFonts
                ea = rFonts.get(qn("w:eastAsia"))
                if ea:
                    return ea
                ascii_f = rFonts.get(qn("w:ascii"))
                if ascii_f:
                    return ascii_f
        except Exception:
            pass
        return run.font.name if run.font.name else None

    # ---- paragraph-level format ----

    @staticmethod
    def _extract_paragraph_format(paragraph) -> dict:
        """Extract paragraph formatting: alignment, indent, spacing.

        The canonical field name exposed to the rule engine is ``align``
        (matching the schema documented in ``config/schema.md``); ``alignment``
        is retained as a backwards-compatible alias.

        New char-based indent fields (first_line_chars, hanging_chars) are
        read directly from OOXML in 1/100 character units (2 chars == 200).
        line_spacing_rule maps the WD_LINE_SPACING enum to a stable string.
        """
        fmt = paragraph.paragraph_format
        alignment = paragraph.alignment or fmt.alignment
        align = DocxExtractor.ALIGNMENT_MAP.get(alignment)
        try:
            ls_rule_enum = fmt.line_spacing_rule
            line_spacing_rule = DocxExtractor.LINE_SPACING_RULE_MAP.get(ls_rule_enum)
        except Exception:
            line_spacing_rule = None
        return {
            "alignment": align,
            "align": align,
            "first_line_indent": DocxExtractor._length_to_cm(fmt.first_line_indent),
            "left_indent": DocxExtractor._length_to_cm(fmt.left_indent),
            "right_indent": DocxExtractor._length_to_cm(fmt.right_indent),
            "line_spacing": DocxExtractor._line_spacing_value(fmt.line_spacing),
            "space_before": DocxExtractor._length_to_pt(fmt.space_before),
            "space_after": DocxExtractor._length_to_pt(fmt.space_after),
            "first_line_chars": DocxExtractor._read_ind_chars(paragraph, "firstLineChars"),
            "hanging_chars": DocxExtractor._read_ind_chars(paragraph, "hangingChars"),
            "line_spacing_rule": line_spacing_rule,
        }

    @staticmethod
    def _read_ind_chars(paragraph, attr: str) -> Optional[int]:
        """Read a char-based indent attribute from w:pPr/w:ind.

        Args:
            paragraph: The paragraph object.
            attr: OOXML attribute local name, e.g. "firstLineChars" or
                "hangingChars".

        Returns:
            int value in 1/100 character units (2 chars == 200), or None.
        """
        try:
            pPr = paragraph._p.pPr
            if pPr is None:
                return None
            ind = pPr.ind
            if ind is None:
                return None
            raw = ind.get(qn(f"w:{attr}"))
            if raw is None:
                return None
            return int(raw)
        except Exception:
            return None

    @staticmethod
    def _length_to_cm(length):
        """Convert EMU length to cm, or None."""
        if length is None:
            return None
        try:
            return round(length.cm, 3)
        except Exception:
            return None

    @staticmethod
    def _length_to_pt(length):
        """Convert EMU length to pt, or None."""
        if length is None:
            return None
        try:
            return round(length.pt, 3)
        except Exception:
            return None

    @staticmethod
    def _line_spacing_value(line_spacing):
        """Parse line-spacing: multiplier (float) or fixed value in pt."""
        if line_spacing is None:
            return None
        try:
            if isinstance(line_spacing, float):
                return round(line_spacing, 3)
            return round(line_spacing.pt, 3)
        except Exception:
            return None

    # ---- tables ----

    def _extract_tables(self, document) -> list:
        """Extract table data as list-of-lists of strings."""
        tables = []
        for table in document.tables:
            rows = []
            for row in table.rows:
                cells = [cell.text for cell in row.cells]
                rows.append(cells)
            tables.append(rows)
        return tables

    # ---- headers / footers ----

    def _extract_headers(self, document) -> list:
        """Extract non-empty header paragraphs from all sections."""
        headers = []
        for section in document.sections:
            for para in section.header.paragraphs:
                if para.text:
                    headers.append(para.text)
        return headers

    def _extract_footers(self, document) -> list:
        """Extract non-empty footer paragraphs from all sections."""
        footers = []
        for section in document.sections:
            for para in section.footer.paragraphs:
                if para.text:
                    footers.append(para.text)
        return footers

    # ---- sections ----

    def _extract_sections(self, document) -> list:
        """Extract section margins and page dimensions in cm."""
        sections = []
        for section in document.sections:
            sections.append({
                "page_width": self._length_to_cm(section.page_width),
                "page_height": self._length_to_cm(section.page_height),
                "margins": {
                    "left": self._length_to_cm(section.left_margin),
                    "right": self._length_to_cm(section.right_margin),
                    "top": self._length_to_cm(section.top_margin),
                    "bottom": self._length_to_cm(section.bottom_margin),
                    "header": self._length_to_cm(section.header_distance),
                    "footer": self._length_to_cm(section.footer_distance),
                },
            })
        return sections
