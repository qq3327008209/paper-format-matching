"""页码 / 页边距 / A4 设置。

从 ``format/engine.py`` 的 ``_apply_page_setup`` 与 ``_create_page_number`` 提取
而来，封装成 ``PageSetupOp``。仅依赖 python-docx 与 ``config``，便于单独测试。
"""
from __future__ import annotations

import logging

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

logger = logging.getLogger(__name__)


_PAGE_NUM_POSITIONS = {
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


class PageSetupOp:
    """Apply page margins, A4 size and footer page numbers."""

    def __init__(self, config) -> None:
        self.config = config

    def apply(self, doc, is_from_txt: bool = False) -> None:
        """Apply page setup to every section of ``doc``."""
        logger.info("正在应用页面边距和页码设置...")

        margin_top = self.config.get("margin_top", 3.7)
        margin_bottom = self.config.get("margin_bottom", 3.5)
        margin_left = self.config.get("margin_left", 2.8)
        margin_right = self.config.get("margin_right", 2.6)
        footer_distance = self.config.get("footer_distance", 1.5)
        should_set_a4 = is_from_txt or self.config.get("force_a4", False)
        add_page_num = self.config.get("page_number", True)
        position = self.config.get("page_number_position", "outside")

        for section in doc.sections:
            section.top_margin = Cm(margin_top)
            section.bottom_margin = Cm(margin_bottom)
            section.left_margin = Cm(margin_left)
            section.right_margin = Cm(margin_right)
            section.footer_distance = Cm(footer_distance)

            if should_set_a4:
                section.page_width = Cm(21)
                section.page_height = Cm(29.7)

            if not add_page_num:
                continue

            if position == "outside":
                doc.settings.odd_and_even_pages_header_footer = True
                section.odd_and_even_pages_header_footer = True

                footer_odd = section.footer
                p_odd = footer_odd.paragraphs[0] if footer_odd.paragraphs else footer_odd.add_paragraph()
                p_odd.clear()
                p_odd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                self._create_page_number(p_odd)

                footer_even = section.even_page_footer
                p_even = footer_even.paragraphs[0] if footer_even.paragraphs else footer_even.add_paragraph()
                p_even.clear()
                p_even.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self._create_page_number(p_even)
            else:
                align = _PAGE_NUM_POSITIONS.get(position, WD_ALIGN_PARAGRAPH.CENTER)
                footer = section.footer
                p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
                p.clear()
                p.alignment = align
                self._create_page_number(p)

        if should_set_a4:
            logger.info("  > 已将页面大小设置为 A4。")

    # ------------------------------------------------------------------
    # Page-number field
    # ------------------------------------------------------------------
    def _create_page_number(self, paragraph) -> None:
        """Append a page-number field to a footer paragraph."""
        font_name = self.config.get("page_number_font", "宋体")
        font_size = self.config.get("page_number_size", 14)
        style = self.config.get("page_number_style", "dash")

        def _field(instruction: str) -> None:
            run = paragraph.add_run()
            _set_run_font(run, font_name, font_size)
            begin = OxmlElement("w:fldChar")
            begin.set(qn("w:fldCharType"), "begin")
            instr_text = OxmlElement("w:instrText")
            instr_text.set(qn("xml:space"), "preserve")
            instr_text.text = instruction
            end = OxmlElement("w:fldChar")
            end.set(qn("w:fldCharType"), "end")
            run._r.extend([begin, instr_text, end])

        if style == "dash":
            run = paragraph.add_run("— ")
            _set_run_font(run, font_name, font_size)
            _field(" PAGE ")
            run = paragraph.add_run(" —")
            _set_run_font(run, font_name, font_size)
        elif style == "page_text":
            run = paragraph.add_run("第 ")
            _set_run_font(run, font_name, font_size)
            _field(" PAGE ")
            run = paragraph.add_run(" 页")
            _set_run_font(run, font_name, font_size)
        elif style == "page_total":
            _field(" PAGE ")
            run = paragraph.add_run(" / ")
            _set_run_font(run, font_name, font_size)
            _field(" NUMPAGES ")
        else:
            _field(" PAGE ")


def _set_run_font(run, font_name: str, font_size: int) -> None:
    """Inline font helper for page-number runs.

    Kept local (not delegated to ParagraphFormatter) because the page-number
    runs are emitted at field-creation time and do not participate in the
    paragraph formatting pipeline.
    """
    from docx.shared import Pt, RGBColor

    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for theme_attr in (
        "w:eastAsiaTheme", "w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:csTheme",
    ):
        rFonts.attrib.pop(qn(theme_attr), None)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    run.font.name = font_name
