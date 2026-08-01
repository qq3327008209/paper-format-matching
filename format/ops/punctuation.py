"""Punctuation normalisation operations.

Copies the half→full-width punctuation replacement, quote pairing and special
pattern protection (URL / email / ISO-standard / time) logic from the legacy
``punctuation.py`` and the Chinese line-break rule logic from the legacy
``east_asian_typography.py`` into this self-contained module. Nothing here
imports from ``scripts`` / ``checker`` / ``wfp``.
"""

from __future__ import annotations

import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from format.ops.space import SpaceOp, redistribute_text_to_runs


# ---------------------------------------------------------------------------
# Quote constants (Unicode escapes for safety)
# ---------------------------------------------------------------------------
LEFT_DOUBLE_QUOTE = "\u201c"   # " 左双引号
RIGHT_DOUBLE_QUOTE = "\u201d"  # " 右双引号
LEFT_SINGLE_QUOTE = "\u2018"   # ' 左单引号
RIGHT_SINGLE_QUOTE = "\u2019"  # ' 右单引号

# Basic half→full-width replacement map (applied only when text has CJK).
REPLACEMENTS = {
    "(": "（",
    ")": "）",
    ":": "：",
    ";": "；",
    "?": "？",
    "!": "！",
}

# Placeholder prefix that cannot appear in normal text.
_PLACEHOLDER_PREFIX = "\x02PROT"

# Double / single quote characters unified into temporary markers for pairing.
_DOUBLE_QUOTE_CHARS = ['"', "\u201c", "\u201d", "\u201e", "\u201f", "\u300c", "\u300d"]
_SINGLE_QUOTE_CHARS = ["'", "\u2018", "\u2019", "\u201a", "\u201b"]


def _protect_special_patterns(text):
    """Protect URL / email / file-path / ISO-standard / time patterns.

    Returns ``(result, protected)`` where ``protected`` is a list of
    ``(placeholder, original)`` pairs restored by :func:`_restore_protected`.
    """
    protected = []
    counter = [0]

    def _replace_with_placeholder(match):
        placeholder = f"{_PLACEHOLDER_PREFIX}{counter[0]}\x03"
        protected.append((placeholder, match.group()))
        counter[0] += 1
        return placeholder

    result = text

    # URLs (http:// https:// ftp://)
    result = re.sub(r"(?:https?|ftp)://\S+", _replace_with_placeholder, result)
    # Email addresses
    result = re.sub(r"[\w.+-]+@[\w-]+\.[\w.-]+", _replace_with_placeholder, result)
    # Windows file paths C:\ D:\
    result = re.sub(r"[A-Za-z]:\\", _replace_with_placeholder, result)
    # Standard numbers (ISO 9001:2015 etc.)
    result = re.sub(r"[A-Za-z]+[\s-]?\d+:\d{2,}", _replace_with_placeholder, result)
    # Time formats HH:MM or HH:MM:SS
    result = re.sub(r"(?<!\d)(\d{1,2}:\d{2}(?::\d{2})?)(?!\d)",
                    _replace_with_placeholder, result)

    return result, protected


def _restore_protected(text, protected):
    """Restore content previously protected by :func:`_protect_special_patterns`."""
    result = text
    for placeholder, original in protected:
        result = result.replace(placeholder, original)
    return result


def has_chinese(text):
    """Return ``True`` when ``text`` contains any CJK character."""
    return bool(re.search(r"[\u4e00-\u9fff]", text))


def fix_text(text):
    """Fix punctuation inside ``text`` (returns a new string)."""
    if not text:
        return text

    result, protected = _protect_special_patterns(text)

    # 1. Ellipsis (before period handling).
    result = re.sub(r"\.{2,}", "……", result)
    result = re.sub(r"。{2,}", "……", result)

    # 2. Em-dash.
    result = re.sub(r"--+", "——", result)
    result = re.sub(r"—(?!—)", "——", result)

    # 3. Basic punctuation replacement (only when CJK present).
    if has_chinese(result):
        for en, cn in REPLACEMENTS.items():
            result = result.replace(en, cn)

    # 4. Comma handling at CJK boundaries.
    result = re.sub(r"([\u4e00-\u9fff]),", r"\1，", result)
    result = re.sub(r",([\u4e00-\u9fff])", r"，\1", result)

    # 5. Period handling at CJK boundaries.
    result = re.sub(r"([\u4e00-\u9fff])\.(\s|$)", r"\1。\2", result)

    # 6. Double-quote pairing.
    temp = result
    for q in _DOUBLE_QUOTE_CHARS:
        temp = temp.replace(q, "\x00")
    if "\x00" in temp:
        chars = list(temp)
        quote_count = 0
        for i, c in enumerate(chars):
            if c == "\x00":
                chars[i] = LEFT_DOUBLE_QUOTE if quote_count % 2 == 0 else RIGHT_DOUBLE_QUOTE
                quote_count += 1
        result = "".join(chars)

    # 7. Single-quote pairing.
    temp = result
    for q in _SINGLE_QUOTE_CHARS:
        temp = temp.replace(q, "\x01")
    if "\x01" in temp:
        chars = list(temp)
        quote_count = 0
        for i, c in enumerate(chars):
            if c == "\x01":
                chars[i] = LEFT_SINGLE_QUOTE if quote_count % 2 == 0 else RIGHT_SINGLE_QUOTE
                quote_count += 1
        result = "".join(chars)

    result = _restore_protected(result, protected)
    return result


def _fix_simple_punctuation(text):
    """Run only the non-pairing replacements (quotes left untouched)."""
    if not text:
        return text

    result, protected = _protect_special_patterns(text)

    result = re.sub(r"\.{2,}", "……", result)
    result = re.sub(r"。{2,}", "……", result)
    result = re.sub(r"--+", "——", result)
    result = re.sub(r"—(?!—)", "——", result)

    if has_chinese(result):
        for en, cn in REPLACEMENTS.items():
            result = result.replace(en, cn)

    result = re.sub(r"([\u4e00-\u9fff]),", r"\1，", result)
    result = re.sub(r",([\u4e00-\u9fff])", r"，\1", result)
    result = re.sub(r"([\u4e00-\u9fff])\.(\s|$)", r"\1。\2", result)

    result = _restore_protected(result, protected)
    return result


def _fix_quotes_whole_text(text):
    """Pair double / single quotes across an entire text block."""
    result = text

    temp = result
    for q in _DOUBLE_QUOTE_CHARS:
        temp = temp.replace(q, "\x00")
    if "\x00" in temp:
        chars = list(temp)
        quote_idx = 0
        for i, c in enumerate(chars):
            if c == "\x00":
                chars[i] = LEFT_DOUBLE_QUOTE if quote_idx % 2 == 0 else RIGHT_DOUBLE_QUOTE
                quote_idx += 1
        result = "".join(chars)

    temp = result
    for q in _SINGLE_QUOTE_CHARS:
        temp = temp.replace(q, "\x01")
    if "\x01" in temp:
        chars = list(temp)
        quote_idx = 0
        for i, c in enumerate(chars):
            if c == "\x01":
                chars[i] = LEFT_SINGLE_QUOTE if quote_idx % 2 == 0 else RIGHT_SINGLE_QUOTE
                quote_idx += 1
        result = "".join(chars)

    return result


# ---------------------------------------------------------------------------
# Chinese line-break rules (copied from east_asian_typography.py)
# ---------------------------------------------------------------------------

_PPR_ORDER = (
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
    "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
    "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
    "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
    "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap",
    "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange",
)
_PPR_ORDER_INDEX = {name: index for index, name in enumerate(_PPR_ORDER)}


def _insert_paragraph_property(p_pr, element, local_name):
    """Insert a property in the order expected by strict OOXML consumers."""
    property_index = _PPR_ORDER_INDEX[local_name]
    for index, child in enumerate(p_pr):
        child_name = child.tag.rsplit("}", 1)[-1]
        if _PPR_ORDER_INDEX.get(child_name, -1) > property_index:
            p_pr.insert(index, element)
            return
    p_pr.append(element)


def _set_paragraph_boolean_property(paragraph, local_name, value):
    p_pr = paragraph._p.get_or_add_pPr()
    tag = qn(f"w:{local_name}")
    element = p_pr.find(tag)
    if element is None:
        element = OxmlElement(f"w:{local_name}")
        _insert_paragraph_property(p_pr, element, local_name)
    desired_value = "1" if value else "0"
    if element.get(qn("w:val")) == desired_value:
        return False
    element.set(qn("w:val"), desired_value)
    return True


def apply_chinese_line_break_rules_to_paragraph(paragraph):
    """Enable Chinese kinsoku and disable hanging punctuation for one paragraph."""
    if not paragraph.text.strip():
        return False
    changed = _set_paragraph_boolean_property(paragraph, "kinsoku", True)
    return _set_paragraph_boolean_property(paragraph, "overflowPunct", False) or changed


def _iter_container_paragraphs(container, seen):
    for paragraph in container.paragraphs:
        paragraph_id = id(paragraph._p)
        if paragraph_id not in seen:
            seen.add(paragraph_id)
            yield paragraph
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_container_paragraphs(cell, seen)


def apply_chinese_line_break_rules(document):
    """Apply Chinese line-break protection across the whole document story."""
    seen = set()
    changed_count = sum(
        apply_chinese_line_break_rules_to_paragraph(paragraph)
        for paragraph in _iter_container_paragraphs(document, seen)
    )
    story_names = (
        "header", "first_page_header", "even_page_header",
        "footer", "first_page_footer", "even_page_footer",
    )
    for section in document.sections:
        for story_name in story_names:
            story = getattr(section, story_name)
            if story.is_linked_to_previous:
                continue
            changed_count += sum(
                apply_chinese_line_break_rules_to_paragraph(paragraph)
                for paragraph in _iter_container_paragraphs(story, seen)
            )
    return changed_count


# ---------------------------------------------------------------------------
# PunctuationOp
# ---------------------------------------------------------------------------

class PunctuationOp:
    """Fix punctuation in text, paragraphs and whole documents."""

    def __init__(self, config) -> None:
        self.config = config
        # Reuse the shared space normaliser for the space step.
        self._space_op = SpaceOp(config)

    def fix_text(self, text):
        """Return ``text`` with punctuation normalised."""
        return fix_text(text)

    def process_paragraph(self, para, space_mode="remove_all"):
        """Fix punctuation in a paragraph's runs in place.

        Simple replacements run per-run (preserving run formatting); quote
        pairing runs across the whole paragraph text. Finally spaces are
        normalised according to ``space_mode``.

        Returns ``True`` when the paragraph was modified.
        """
        full_text = para.text
        if not full_text.strip():
            return False

        runs = para.runs
        if not runs:
            return False

        changed = False

        # 1. Per-run simple punctuation replacement.
        for run in runs:
            original = run.text
            fixed = _fix_simple_punctuation(original)
            if fixed != original:
                run.text = fixed
                changed = True

        # 2. Quote pairing across the whole paragraph.
        full_after_simple = para.text
        full_after_quotes = _fix_quotes_whole_text(full_after_simple)
        if full_after_quotes != full_after_simple:
            redistribute_text_to_runs(runs, full_after_quotes)
            changed = True

        # 3. Space normalisation.
        if self._space_op.process_paragraph(para, space_mode):
            changed = True

        return changed

    def process_document(self, doc):
        """Fix punctuation in every paragraph (including table cells).

        Returns the number of paragraphs that were modified.
        """
        changed_count = 0

        for para in doc.paragraphs:
            if self.process_paragraph(para):
                changed_count += 1

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if self.process_paragraph(para):
                            changed_count += 1

        # Apply Chinese kinsoku / line-break rules document-wide.
        apply_chinese_line_break_rules(doc)

        return changed_count
