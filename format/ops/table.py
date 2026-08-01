"""Table formatting operations.

Applies font / size / row-height / borders / header-bold (and table line
spacing) to every cell of a table, driven by the engine ``Config``. Refactored
from the legacy ``_format_tables`` logic.
"""

from __future__ import annotations

import re

from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from format.ops.font import FontOp


# Numeric table-text patterns used for smart alignment (copied from wfp_core).
RE_CURRENCY_PREFIX = re.compile(r"^[¥￥$]")
RE_CURRENCY_SUFFIX = re.compile(r"(元|万元|亿元)$")
RE_NUMERIC_TABLE_TEXT = re.compile(r"^[-+]?(?:\d+(?:\.\d+)?|\.\d+)%?$")


def _get_or_add_table_pr(table):
    """Return (creating if needed) the ``<w:tblPr>`` element of ``table``."""
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    return tbl_pr


def _set_table_borders(table, size_pt=0.5, color="000000"):
    """Apply unified single-line borders to the whole table."""
    size = max(1, int(float(size_pt) * 8))
    tbl_pr = _get_or_add_table_pr(table)
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    else:
        for child in list(borders):
            borders.remove(child)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), str(size))
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)
        borders.append(elem)


def _set_cell_borders(cell, size_pt=0.5, color="000000"):
    """Apply single-line borders to a single cell."""
    size = max(1, int(float(size_pt) * 8))
    tc = cell._tc
    tc_pr = tc.tcPr
    if tc_pr is None:
        tc_pr = OxmlElement("w:tcPr")
        tc.insert(0, tc_pr)

    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    else:
        for child in list(borders):
            borders.remove(child)

    for edge in ("top", "left", "bottom", "right"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), str(size))
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)
        borders.append(elem)


def _is_numeric_table_text(text):
    text = (text or "").strip()
    if not text:
        return False
    text = text.replace(",", "").replace("，", "").replace("％", "%")
    text = RE_CURRENCY_PREFIX.sub("", text)
    text = RE_CURRENCY_SUFFIX.sub("", text)
    return RE_NUMERIC_TABLE_TEXT.match(text) is not None


def _to_float(value, default=0.0):
    try:
        if value == "" or value is None:
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


class TableOp:
    """Format tables: fonts, sizes, row heights, borders and header bold."""

    def __init__(self, config) -> None:
        self.config = config
        self._font_op = FontOp(config)

    def format_table(self, table, apply_color: bool = True) -> None:
        """Apply configured font / size / row-height / borders / header-bold."""
        table_font = self.config.get(
            "table_font", self.config.get("body_font", "仿宋_GB2312")
        )
        table_header_font = self.config.get("table_header_font", table_font)
        font_en = self.config.get("body_font_en", "Times New Roman")
        table_size = _to_float(
            self.config.get("table_size", self.config.get("body_size", 12)), 12.0
        )
        table_line_spacing = _to_float(self.config.get("table_line_spacing", 22), 22.0)
        row_height_cm = _to_float(self.config.get("table_row_height_cm", 0.7), 0.7)
        border_size_pt = _to_float(self.config.get("table_border_size_pt", 0.5), 0.5)
        unified_borders = self.config.get("table_unified_borders", True)
        header_bold = self.config.get("table_header_bold", True)
        smart_align = self.config.get("table_smart_align", False)

        if unified_borders:
            _set_table_borders(table, size_pt=border_size_pt)

        # Detect a "序号" (serial-number) column for smart alignment.
        serial_col_idx = None
        if table.rows:
            for col_idx, cell in enumerate(table.rows[0].cells):
                head_text = "".join(p.text for p in cell.paragraphs).strip()
                if "序号" in head_text or head_text == "序":
                    serial_col_idx = col_idx
                    break

        for row_idx, row in enumerate(table.rows):
            if row_height_cm > 0:
                row.height = Cm(row_height_cm)
                row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

            for col_idx, cell in enumerate(row.cells):
                if unified_borders:
                    _set_cell_borders(cell, size_pt=border_size_pt)

                cell_text = "".join(p.text for p in cell.paragraphs).strip()
                is_header_row = row_idx == 0
                font_name = table_header_font if is_header_row else table_font

                for para in cell.paragraphs:
                    if para.text.strip():
                        for run in para.runs:
                            self._font_op.set_font(
                                run, font_name, font_en, table_size,
                                set_color=apply_color,
                            )
                            if is_header_row and header_bold:
                                run.font.bold = True

                    para.paragraph_format.first_line_indent = Pt(0)
                    para.paragraph_format.space_before = Pt(0)
                    para.paragraph_format.space_after = Pt(0)
                    if table_line_spacing > 0:
                        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                        para.paragraph_format.line_spacing = Pt(table_line_spacing)
                    else:
                        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

                    if smart_align:
                        from docx.enum.text import WD_ALIGN_PARAGRAPH
                        if is_header_row:
                            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        elif "合计" in cell_text or "总计" in cell_text:
                            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        elif serial_col_idx is not None and col_idx == serial_col_idx:
                            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        elif _is_numeric_table_text(cell_text):
                            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        else:
                            para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def format_all(self, doc, apply_color: bool = True) -> None:
        """Format every table in ``doc``."""
        if not self.config.get("enable_table_formatting", False):
            return
        for table in doc.tables:
            self.format_table(table, apply_color=apply_color)
