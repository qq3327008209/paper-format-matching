"""Paragraph classification and per-paragraph formatting.

``ParagraphClassifier`` owns the large ``apply_para_format`` method that
decides whether a paragraph is a heading (h1 / h2 / h3 / h4), an attachment
marker, a picture-text paragraph, or ordinary body text -- and then applies
the configured font, size, alignment, indent, outline level and pagination
properties accordingly.

It is a focused extraction of the classification logic that previously lived
inside ``ParagraphFormatter``.
"""

from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from format.ops.font import FontOp
from format.ops.heading import (
    RE_ATTACHMENT,
    RE_H2_INLINE_TITLE,
    RE_HEADING_H1,
    RE_HEADING_H2,
    RE_HEADING_H3,
    RE_HEADING_H4,
)
from format.ops.paragraph import ParagraphOp
from format.ops.paragraph_formatter import (
    apply_font_to_runs,
    apply_text_indent_and_align,
    format_heading,
    get_paragraph_alignment,
    reset_pagination_properties,
    set_run_font,
    strip_leading_whitespace,
)
from format.ops.symbol_normalizer import (
    has_drawing_or_pict,
    has_embedded_object,
)


class ParagraphClassifier:
    """Classify a single paragraph and apply the configured formatting."""

    def __init__(self, config, logger) -> None:
        """Initialise the classifier.

        Args:
            config: A ``format.config.Config`` instance.
            logger: A ``logging.Logger`` used for progress messages.
        """
        self.config = config
        self._logger = logger
        self._font_op = FontOp(config)
        self._heading_op = self._make_heading_op(config)
        self._para_op = ParagraphOp(config)

    def _make_heading_op(self, config):
        """Return a HeadingOp instance.

        Indirected through a helper so tests can patch it if needed.
        """
        from format.ops.heading import HeadingOp
        return HeadingOp(config)

    def _log(self, message: str) -> None:
        self._logger.info(message)

    def apply_para_format(
        self,
        block,
        font_op=None,
        para_op=None,
        is_from_txt: bool = False,
        apply_color: bool = False,
    ) -> None:
        """Classify and format a single paragraph block.

        ``font_op`` / ``para_op`` are accepted to match the Pipeline contract
        but the actual formatting is applied directly via this classifier's
        helpers (mirroring the legacy behaviour).
        """
        if not isinstance(block, Paragraph):
            return
        para = block
        if not para.text.strip():
            return

        body_font = self.config.get('body_font', '仿宋_GB2312')
        body_size = self.config.get('body_size', 16)
        h1_font = self.config.get('h1_font', '黑体')
        h1_size = self.config.get('h1_size', 16)
        h2_font = self.config.get('h2_font', '楷体_GB2312')
        h2_size = self.config.get('h2_size', 16)
        attachment_font = self.config.get('attachment_font', '仿宋_GB2312')
        attachment_size = self.config.get('attachment_size', 16)

        original_text = para.text
        text_to_check = para.text.lstrip()
        text_stripped = para.text.strip()
        leading_space_count = len(original_text) - len(text_to_check)
        para_text_preview = text_to_check[:30].replace("\n", " ")

        # Base line spacing for the paragraph
        spacing = para._p.get_or_add_pPr().get_or_add_spacing()
        spacing.set(qn('w:beforeAutospacing'), '0')
        spacing.set(qn('w:afterAutospacing'), '0')
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        line_spacing = self.config.get('line_spacing', 28) or 28
        para.paragraph_format.line_spacing = Pt(line_spacing)

        # Paragraphs that carry pictures / embedded objects: format text only
        is_pic = has_drawing_or_pict(para)
        is_embedded_obj = has_embedded_object(para)
        if is_pic or is_embedded_obj:
            log_msg = "图片" if is_pic else "嵌入对象"
            self._log(f"段落: {log_msg} - 仅格式化文字")
            if RE_HEADING_H1.match(text_to_check):
                self._log(f"  > 文字识别为一级标题: \"{para_text_preview}...\"")
                apply_font_to_runs(
                    para, h1_font, h1_size, self._font_op, self.config,
                    set_color=apply_color,
                )
            elif RE_HEADING_H2.match(text_to_check):
                self._log(f"  > 文字识别为二级标题: \"{para_text_preview}...\"")
                apply_font_to_runs(
                    para, h2_font, h2_size, self._font_op, self.config,
                    set_color=apply_color,
                )
            elif RE_HEADING_H3.match(text_to_check):
                self._log(f"  > 文字识别为三级标题: \"{para_text_preview}...\"")
                apply_font_to_runs(
                    para, body_font, body_size, self._font_op, self.config,
                    set_color=apply_color,
                )
            elif RE_HEADING_H4.match(text_to_check):
                self._log(f"  > 文字识别为四级标题: \"{para_text_preview}...\"")
                apply_font_to_runs(
                    para, body_font, body_size, self._font_op, self.config,
                    set_color=apply_color,
                )
            elif text_to_check:
                self._log(f"  > 文字识别为正文: \"{para_text_preview}...\"")
                apply_font_to_runs(
                    para, body_font, body_size, self._font_op, self.config,
                    set_color=apply_color,
                )
            return

        # Attachment marker
        is_attachment_enabled = self.config.get('enable_attachment_formatting', True)
        is_attachment_candidate = False
        if is_from_txt:
            if RE_ATTACHMENT.match(text_stripped):
                is_attachment_candidate = True
        else:
            align = get_paragraph_alignment(para)
            if align in (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.JUSTIFY, None) and RE_ATTACHMENT.match(text_stripped):
                is_attachment_candidate = True

        if is_attachment_enabled and is_attachment_candidate:
            self._log(f"段落: 附件标识 - \"{para_text_preview}...\"")
            strip_leading_whitespace(para, self._logger)
            apply_font_to_runs(
                para, attachment_font, attachment_size, self._font_op,
                self.config, set_color=apply_color,
            )
            reset_pagination_properties(para)
            para.paragraph_format.page_break_before = True
            para.paragraph_format.left_indent = Pt(0)
            para.paragraph_format.first_line_indent = None
            ind = para._p.get_or_add_pPr().get_or_add_ind()
            ind.set(qn("w:firstLineChars"), "0")
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            format_heading(para, 1, self.config, self._logger)
            return

        if RE_HEADING_H1.match(text_to_check):
            self._log(f"段落: 一级标题 - \"{para_text_preview}...\"")
            strip_leading_whitespace(para, self._logger)
            format_heading(para, 1, self.config, self._logger)
            apply_font_to_runs(
                para, h1_font, h1_size, self._font_op, self.config,
                set_color=apply_color,
            )
            apply_text_indent_and_align(para, self.config)
            reset_pagination_properties(para)

        elif RE_HEADING_H2.match(text_to_check):
            self._log(f"段落: 二级标题 - \"{para_text_preview}...\"")
            strip_leading_whitespace(para, self._logger)

            parts = para.text.split('。', 1)

            if len(parts) == 2 and parts[1].strip():
                self._log("  > 检测到二级标题与正文在同一段落，执行段内格式拆分。")
                title_len = len(parts[0]) + 1

                original_runs = []
                for r in para.runs:
                    original_runs.append({
                        'text': r.text, 'bold': r.bold, 'italic': r.italic,
                        'underline': r.underline, 'font_color': r.font.color.rgb,
                    })

                para.clear()

                char_count = 0
                for run_info in original_runs:
                    run_text = run_info['text']
                    run_end_pos = char_count + len(run_text)

                    title_run, body_run, new_run = None, None, None

                    if run_end_pos <= title_len:
                        new_run = para.add_run(run_text)
                        set_run_font(
                            new_run, h2_font, h2_size, self._font_op,
                            self.config, set_color=apply_color,
                        )

                    elif char_count >= title_len:
                        new_run = para.add_run(run_text)
                        set_run_font(
                            new_run, body_font, body_size, self._font_op,
                            self.config, set_color=apply_color,
                        )

                    else:
                        split_index = title_len - char_count
                        title_part = run_text[:split_index]
                        body_part = run_text[split_index:]

                        if title_part:
                            title_run = para.add_run(title_part)
                            set_run_font(
                                title_run, h2_font, h2_size, self._font_op,
                                self.config, set_color=apply_color,
                            )
                        if body_part:
                            body_run = para.add_run(body_part)
                            set_run_font(
                                body_run, body_font, body_size, self._font_op,
                                self.config, set_color=apply_color,
                            )

                    runs_to_format = (
                        [r for r in [title_run, body_run] if r]
                        or ([new_run] if new_run else [])
                    )
                    for r in runs_to_format:
                        if r:
                            r.bold = run_info['bold']
                            r.italic = run_info['italic']
                            r.underline = run_info['underline']
                            if run_info['font_color']:
                                r.font.color.rgb = run_info['font_color']

                    char_count = run_end_pos

                format_heading(para, 2, self.config, self._logger)
                apply_text_indent_and_align(para, self.config)
                reset_pagination_properties(para)

            else:
                match = RE_H2_INLINE_TITLE.match(text_to_check)
                if match and not (text_to_check.startswith('（') and text_to_check.strip().endswith('）')):
                    self._log("  > 已将二级标题的括号统一为中文括号。")
                    for r in para.runs:
                        r.text = r.text.replace('(', '（', 1).replace(')', '）', 1)
                format_heading(para, 2, self.config, self._logger)
                apply_font_to_runs(
                    para, h2_font, h2_size, self._font_op, self.config,
                    set_color=apply_color,
                )
                apply_text_indent_and_align(para, self.config)
                reset_pagination_properties(para)

        elif RE_HEADING_H3.match(text_to_check):
            self._log(f"段落: 三级标题 - \"{para_text_preview}...\"")
            strip_leading_whitespace(para, self._logger)
            format_heading(para, 3, self.config, self._logger)
            apply_font_to_runs(
                para, body_font, body_size, self._font_op, self.config,
                set_color=apply_color,
            )
            apply_text_indent_and_align(para, self.config)
            reset_pagination_properties(para)

        elif RE_HEADING_H4.match(text_to_check):
            self._log(f"段落: 四级标题 - \"{para_text_preview}...\"")
            strip_leading_whitespace(para, self._logger)
            format_heading(para, 4, self.config, self._logger)
            apply_font_to_runs(
                para, body_font, body_size, self._font_op, self.config,
                set_color=apply_color,
            )
            apply_text_indent_and_align(para, self.config)
            reset_pagination_properties(para)

        elif not is_from_txt:
            para_alignment = get_paragraph_alignment(para)
            if para_alignment in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT):
                align_text = "居中" if para_alignment == WD_ALIGN_PARAGRAPH.CENTER else "右对齐"
                self._log(f"段落: {align_text}正文 - 保留原对齐")
                apply_font_to_runs(
                    para, body_font, body_size, self._font_op, self.config,
                    set_color=apply_color,
                )
                reset_pagination_properties(para)
            elif leading_space_count > 5:
                self._log(f"段落: 正文 (保留前导空格) - \"{para_text_preview}...\"")
                apply_font_to_runs(
                    para, body_font, body_size, self._font_op, self.config,
                    set_color=apply_color,
                )
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                reset_pagination_properties(para)
            elif (para.paragraph_format.first_line_indent is None or para.paragraph_format.first_line_indent.pt == 0) and leading_space_count == 0:
                self._log(f"段落: 正文 (保留0缩进) - \"{para_text_preview}...\"")
                apply_font_to_runs(
                    para, body_font, body_size, self._font_op, self.config,
                    set_color=apply_color,
                )
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                reset_pagination_properties(para)
            else:
                self._log(f"段落: 正文 (应用标准缩进) - \"{para_text_preview}...\"")
                strip_leading_whitespace(para, self._logger)
                apply_font_to_runs(
                    para, body_font, body_size, self._font_op, self.config,
                    set_color=apply_color,
                )
                apply_text_indent_and_align(para, self.config)
                reset_pagination_properties(para)
        else:
            self._log(f"段落: 正文 (源自TXT，强制缩进) - \"{para_text_preview}...\"")
            strip_leading_whitespace(para, self._logger)
            apply_font_to_runs(
                para, body_font, body_size, self._font_op, self.config,
                set_color=apply_color,
            )
            apply_text_indent_and_align(para, self.config)
            reset_pagination_properties(para)
