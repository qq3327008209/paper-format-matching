# -*- coding: utf-8 -*-
"""Core formatting engine — thin facade.

``FormatEngine`` is the public entry point that :class:`format.pipeline.Pipeline`
drives.  The heavy lifting has been split into focused sub-modules under
``format.ops``:

- :class:`format.ops.text_loader.TextLoader` — file reading / Markdown cleanup
- :class:`format.ops.symbol_normalizer.SymbolNormalizer` — Chinese punctuation
- :class:`format.ops.paragraph_formatter.ParagraphFormatter` — paragraph
  classification, font/size/alignment, title detection
- :class:`format.ops.table.TableOp` — table borders / widths / fonts

This module re-exports the regex constants and the legacy static-method names
so existing call sites (``format.engine.RE_*``, ``engine._normalize_ellipsis``)
keep working without modification.
"""

import logging
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.shared import Cm, Pt

from format.converters.office import convert_to_docx
from format.ops.heading import (
    CHINESE_NUM_PATTERN,
    RE_ATTACHMENT,
    RE_HEADING_H1,
    RE_HEADING_H2,
    RE_HEADING_H3,
    RE_HEADING_H4,
    RE_H2_INLINE_TITLE,
    RE_TITLE_H1,
    RE_TITLE_H2,
)
from format.ops.paragraph_formatter import ParagraphFormatter
from format.ops.symbol_normalizer import (
    SymbolNormalizer,
    has_chinese,
    has_drawing_or_pict,
    has_embedded_object,
    has_field_codes,
    is_after_digit_or_latin,
    is_before_digit_or_latin,
)
from format.ops.table import TableOp
from format.ops.text_loader import TextLoader

logger = logging.getLogger(__name__)


class FormatEngine:
    """Core formatting engine.

    Owns a :class:`format.config.Config` instance and exposes the methods that
    :class:`format.pipeline.Pipeline` drives to format a document.

    The implementation is delegated to focused sub-modules; this class is a
    thin facade that wires them together.
    """

    def __init__(self, config):
        """Initialize the engine.

        Args:
            config: A :class:`format.config.Config` instance providing all
                formatting parameters via ``get``/``get_nested``/``merged``.
        """
        self.config = config
        self._text_loader = TextLoader()
        self._symbol_normalizer = SymbolNormalizer()
        self._paragraph_formatter = ParagraphFormatter(config, logger)
        self._table_op = TableOp(config)

    # ------------------------------------------------------------------
    # Logging helper
    # ------------------------------------------------------------------
    def _log(self, message):
        logger.info(message)

    # ------------------------------------------------------------------
    # Public entry point
    # ------------------------------------------------------------------
    def run(self, input_path, output_path):
        """Run the full formatting pipeline.

        ``Pipeline`` is imported lazily to avoid a circular import
        (``pipeline`` imports ``FormatEngine`` for type-checking).
        """
        from .pipeline import Pipeline
        return Pipeline(self).run(input_path, output_path)

    # ------------------------------------------------------------------
    # Document loading
    # ------------------------------------------------------------------
    def load_document(self, input_path):
        """Load a source document into a python-docx ``Document``.

        ``.doc``/``.wps``/``.docx`` files are converted to a temporary
        ``.docx`` via :func:`format.converters.office.convert_to_docx`.
        ``.txt``/``.md`` files are read, cleaned and turned into an in-memory
        ``Document``.
        """
        file_ext = os.path.splitext(input_path)[1].lower()

        if file_ext in ('.docx', '.doc', '.wps'):
            self._log("正在加载并转换文档为 .docx ...")
            docx_path = convert_to_docx(input_path)
            return Document(docx_path)

        if file_ext == '.txt':
            self._log("检测到 .txt 文件，正在创建 .docx ...")
            text_content = self._text_loader.read_text_file(input_path)
            text_content = self._normalize_text_blank_lines(text_content)
            doc = Document()
            for line in text_content.split('\n'):
                doc.add_paragraph(line.strip())
            self._log("TXT 转换完成。")
            return doc

        if file_ext == '.md':
            self._log("检测到 .md 文件，正在清理 Markdown 标记并创建 .docx ...")
            raw_text = self._text_loader.read_text_file(input_path)
            cleaned_text = self._text_loader.clean_markdown(raw_text)
            cleaned_text = self._normalize_text_blank_lines(cleaned_text)
            doc = Document()
            for line in cleaned_text.split('\n'):
                doc.add_paragraph(line.strip())
            self._log("Markdown 转换完成。")
            return doc

        raise ValueError(f"不支持的文件格式: {file_ext}")

    def _file_is_text(self, input_path):
        """Return True when the source is a plain-text format (.txt/.md)."""
        return os.path.splitext(input_path)[1].lower() in ('.txt', '.md')

    # ------------------------------------------------------------------
    # Blank-line normalization
    # ------------------------------------------------------------------
    def _normalize_text_blank_lines(self, text, keep_single_blank_lines=False):
        mode = self.config.get('blank_line_mode', 'delete_single')
        if mode in ('preserve', 'none', '不改动任何空行'):
            return text
        keep_single = mode in (
            'keep_single', 'compress', '保留单个空行，多个空行保留至1个空行',
        )
        if not self.config.get('remove_blank_lines', True):
            keep_single = True
        return self._text_loader.remove_blank_lines_from_text(
            text, keep_single_blank_lines=keep_single,
        )

    # ------------------------------------------------------------------
    # Chinese symbol / punctuation normalization
    # (delegated to SymbolNormalizer; legacy method names preserved for
    #  backwards compatibility and tests)
    # ------------------------------------------------------------------
    @staticmethod
    def _paragraph_has_ooxml(para, tag_name):
        return SymbolNormalizer.paragraph_has_ooxml(para, tag_name)

    @classmethod
    def _has_field_codes(cls, para):
        return has_field_codes(para)

    @classmethod
    def _has_drawing_or_pict(cls, para):
        return has_drawing_or_pict(para)

    @classmethod
    def _has_embedded_object(cls, para):
        return has_embedded_object(para)

    def _has_drawing(self, block):
        """Return True when a paragraph block carries a drawing/picture."""
        from docx.text.paragraph import Paragraph
        if isinstance(block, Paragraph):
            return has_drawing_or_pict(block)
        return False

    @staticmethod
    def _has_chinese(text):
        return has_chinese(text)

    @staticmethod
    def _is_digit_or_latin(ch):
        return SymbolNormalizer.is_digit_or_latin(ch)

    @staticmethod
    def _prev_non_space(text, index):
        return SymbolNormalizer.prev_non_space(text, index)

    @staticmethod
    def _next_non_space(text, index):
        return SymbolNormalizer.next_non_space(text, index)

    @classmethod
    def _is_after_digit_or_latin(cls, text, index):
        return is_after_digit_or_latin(text, index)

    @classmethod
    def _is_before_digit_or_latin(cls, text, index):
        return is_before_digit_or_latin(text, index)

    @staticmethod
    def _normalize_ellipsis(text):
        return SymbolNormalizer.normalize_ellipsis(text)

    @classmethod
    def _normalize_simple_punctuation(cls, text):
        return SymbolNormalizer.normalize_simple_punctuation(text)

    @staticmethod
    def _normalize_bracket_pairs(text):
        return SymbolNormalizer.normalize_bracket_pairs(text)

    @staticmethod
    def _normalize_quote_pairs(
        text, quote_chars, left_quote, right_quote, skip_inner_latin=False,
    ):
        return SymbolNormalizer.normalize_quote_pairs(
            text, quote_chars, left_quote, right_quote,
            skip_inner_latin=skip_inner_latin,
        )

    @classmethod
    def _normalize_symbols_in_text(cls, text):
        return SymbolNormalizer.normalize_symbols_in_text(text)

    @staticmethod
    def _redistribute_text_to_runs(runs, new_full_text):
        from format.ops.space import redistribute_text_to_runs as _redist
        _redist(runs, new_full_text)

    def _normalize_paragraph_symbols(self, para):
        return self._symbol_normalizer.normalize_paragraph_symbols(para)

    def _normalize_document_symbols(self, doc):
        """Normalize Chinese punctuation across paragraphs and table cells."""
        return self._symbol_normalizer.normalize_document_symbols(doc)

    # ------------------------------------------------------------------
    # Block iteration
    # ------------------------------------------------------------------
    def _iter_block_items(self, parent):
        """Yield Paragraph and Table objects in document order."""
        return self._paragraph_formatter._iter_block_items(parent)

    # ------------------------------------------------------------------
    # Paragraph formatting
    # ------------------------------------------------------------------
    def _set_run_font(self, run, font_name, size_pt, set_color=False):
        self._paragraph_formatter._set_run_font(run, font_name, size_pt, set_color=set_color)

    def _apply_font_to_runs(self, para, font_name, size_pt, set_color=False):
        self._paragraph_formatter._apply_font_to_runs(para, font_name, size_pt, set_color=set_color)

    def _get_paragraph_font_info(self, para):
        return self._paragraph_formatter._get_paragraph_font_info(para)

    @staticmethod
    def _get_paragraph_alignment(para):
        return ParagraphFormatter._get_paragraph_alignment(para)

    def _strip_leading_whitespace(self, para):
        self._paragraph_formatter._strip_leading_whitespace(para)

    def _reset_pagination_properties(self, para):
        self._paragraph_formatter._reset_pagination_properties(para)

    def _get_outline_level(self, para):
        return self._paragraph_formatter._get_outline_level(para)

    def _set_outline_level(self, para, level):
        return self._paragraph_formatter._set_outline_level(para, level)

    def _format_heading(self, para, level):
        self._paragraph_formatter._format_heading(para, level)

    def _apply_text_indent_and_align(self, para):
        self._paragraph_formatter._apply_text_indent_and_align(para)

    def _find_title_paragraphs(self, doc, is_from_txt, start_index=0):
        """Find the main title and subtitle paragraph block indices."""
        return self._paragraph_formatter.find_title_paragraphs(
            doc, is_from_txt, start_index=start_index,
        )

    def _apply_para_format(self, block, font_op, para_op, is_from_txt, apply_color):
        """Classify and format a single paragraph block."""
        self._paragraph_formatter.apply_para_format(
            block, font_op, para_op, is_from_txt, apply_color,
        )

    # ------------------------------------------------------------------
    # Table formatting
    # ------------------------------------------------------------------
    def _format_tables(self, doc, apply_color=True):
        """Format every table: borders, widths, row height, fonts, alignment."""
        self._table_op.format_all(doc, apply_color=apply_color)

    # ------------------------------------------------------------------
    # Page setup & page numbers
    # ------------------------------------------------------------------
    def _apply_page_setup(self, doc, is_from_txt=False):
        """Apply page margins, A4 size and footer page numbers."""
        self._log("正在应用页面边距和页码设置...")

        from docx.enum.text import WD_ALIGN_PARAGRAPH

        margin_top = self.config.get('margin_top', 3.7)
        margin_bottom = self.config.get('margin_bottom', 3.5)
        margin_left = self.config.get('margin_left', 2.8)
        margin_right = self.config.get('margin_right', 2.6)
        footer_distance = self.config.get('footer_distance', 1.5)
        should_set_a4 = is_from_txt or self.config.get('force_a4', False)
        add_page_num = self.config.get('page_number', True)
        position = self.config.get('page_number_position', 'outside')

        for section in doc.sections:
            section.top_margin = Cm(margin_top)
            section.bottom_margin = Cm(margin_bottom)
            section.left_margin = Cm(margin_left)
            section.right_margin = Cm(margin_right)
            section.footer_distance = Cm(footer_distance)

            if should_set_a4:
                section.page_width = Cm(21)
                section.page_height = Cm(29.7)

            if not add_page_num:
                continue

            if position == 'outside':
                doc.settings.odd_and_even_pages_header_footer = True
                section.odd_and_even_pages_header_footer = True

                footer_odd = section.footer
                p_odd = footer_odd.paragraphs[0] if footer_odd.paragraphs else footer_odd.add_paragraph()
                p_odd.clear()
                p_odd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                self._create_page_number(p_odd)

                footer_even = section.even_page_footer
                p_even = footer_even.paragraphs[0] if footer_even.paragraphs else footer_even.add_paragraph()
                p_even.clear()
                p_even.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self._create_page_number(p_even)
            else:
                align_map = {
                    'center': WD_ALIGN_PARAGRAPH.CENTER,
                    'left': WD_ALIGN_PARAGRAPH.LEFT,
                    'right': WD_ALIGN_PARAGRAPH.RIGHT,
                }
                align = align_map.get(position, WD_ALIGN_PARAGRAPH.CENTER)
                footer = section.footer
                p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
                p.clear()
                p.alignment = align
                self._create_page_number(p)

        if should_set_a4:
            self._log("  > 已将页面大小设置为 A4。")

    def _create_page_number(self, paragraph):
        """Append a page-number field to a footer paragraph."""
        font_name = self.config.get('page_number_font', '宋体')
        font_size = self.config.get('page_number_size', 14)
        style = self.config.get('page_number_style', 'dash')

        def _field(instruction):
            run = paragraph.add_run()
            self._paragraph_formatter._set_run_font(
                run, font_name, font_size, set_color=True,
            )
            begin = OxmlElement('w:fldChar')
            begin.set(qn('w:fldCharType'), 'begin')
            instr_text = OxmlElement('w:instrText')
            instr_text.set(qn('xml:space'), 'preserve')
            instr_text.text = instruction
            end = OxmlElement('w:fldChar')
            end.set(qn('w:fldCharType'), 'end')
            run._r.extend([begin, instr_text, end])

        if style == 'dash':
            run = paragraph.add_run('— ')
            self._paragraph_formatter._set_run_font(
                run, font_name, font_size, set_color=True,
            )
            _field(' PAGE ')
            run = paragraph.add_run(' —')
            self._paragraph_formatter._set_run_font(
                run, font_name, font_size, set_color=True,
            )
        elif style == 'page_text':
            run = paragraph.add_run('第 ')
            self._paragraph_formatter._set_run_font(
                run, font_name, font_size, set_color=True,
            )
            _field(' PAGE ')
            run = paragraph.add_run(' 页')
            self._paragraph_formatter._set_run_font(
                run, font_name, font_size, set_color=True,
            )
        elif style == 'page_total':
            _field(' PAGE ')
            run = paragraph.add_run(' / ')
            self._paragraph_formatter._set_run_font(
                run, font_name, font_size, set_color=True,
            )
            _field(' NUMPAGES ')
        else:
            _field(' PAGE ')
