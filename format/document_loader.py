"""Document loading: convert .doc/.wps/.txt/.md/.docx into python-docx Document.

Refactored from ``FormatEngine.load_document`` + ``_file_is_text`` +
``_normalize_text_blank_lines``. Owns its own ``TextLoader`` for text-file
preparation; delegates .doc/.wps → .docx to :mod:`format.converters.office`.
"""
from __future__ import annotations

import logging
import os

from docx import Document

from format.converters.office import convert_to_docx
from format.ops.text_loader import TextLoader

logger = logging.getLogger(__name__)


_TEXT_EXTS = (".txt", ".md")
_OFFICE_EXTS = (".docx", ".doc", ".wps")


class DocumentLoader:
    """Load any supported source file into a python-docx ``Document``."""

    def __init__(self, config) -> None:
        self.config = config
        self._text_loader = TextLoader()

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------
    def load(self, input_path: str) -> Document:
        """Load a source document into a python-docx ``Document``."""
        file_ext = os.path.splitext(input_path)[1].lower()

        if file_ext in _OFFICE_EXTS:
            logger.info("正在加载并转换文档为 .docx ...")
            docx_path = convert_to_docx(input_path)
            return Document(docx_path)

        if file_ext == ".txt":
            logger.info("检测到 .txt 文件，正在创建 .docx ...")
            text_content = self._text_loader.read_text_file(input_path)
            text_content = self._normalize_text_blank_lines(text_content)
            doc = Document()
            for line in text_content.split("\n"):
                doc.add_paragraph(line.strip())
            logger.info("TXT 转换完成。")
            return doc

        if file_ext == ".md":
            logger.info("检测到 .md 文件，正在清理 Markdown 标记并创建 .docx ...")
            raw_text = self._text_loader.read_text_file(input_path)
            cleaned_text = self._text_loader.clean_markdown(raw_text)
            cleaned_text = self._normalize_text_blank_lines(cleaned_text)
            doc = Document()
            for line in cleaned_text.split("\n"):
                doc.add_paragraph(line.strip())
            logger.info("Markdown 转换完成。")
            return doc

        raise ValueError(f"不支持的文件格式: {file_ext}")

    def is_text_input(self, input_path: str) -> bool:
        """Return True when the source is a plain-text format (.txt/.md)."""
        return os.path.splitext(input_path)[1].lower() in _TEXT_EXTS

    # ------------------------------------------------------------------
    # Blank-line normalisation (delegated)
    # ------------------------------------------------------------------
    def _normalize_text_blank_lines(
        self, text: str, keep_single_blank_lines: bool = False,
    ) -> str:
        mode = self.config.get("blank_line_mode", "delete_single")
        if mode in ("preserve", "none", "不改动任何空行"):
            return text
        keep_single = mode in (
            "keep_single", "compress", "保留单个空行，多个空行保留至1个空行",
        )
        if not self.config.get("remove_blank_lines", True):
            keep_single = True
        return self._text_loader.remove_blank_lines_from_text(
            text, keep_single_blank_lines=keep_single,
        )
